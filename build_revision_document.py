from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = "Patient_Consultation_Revision_Technical_Documentation.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.7)
sec.left_margin, sec.right_margin = Inches(.72), Inches(.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].font.color.rgb = RGBColor(33, 41, 52)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.08
for s, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
    styles[s].font.name = 'Aptos Display' if s != 'Normal' else 'Aptos'
    styles[s].font.size = Pt(size)
    styles[s].font.color.rgb = RGBColor(0, 0, 0)
    styles[s].font.bold = True
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Heading 1'].paragraph_format.space_before = Pt(14)
styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].paragraph_format.space_before = Pt(10)
styles['Heading 2'].paragraph_format.space_after = Pt(5)

if 'Code Path' not in styles:
    st = styles.add_style('Code Path', WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = 'Aptos Mono'; st.font.size = Pt(8.7); st.font.color.rgb = RGBColor(28, 53, 83)
    st.paragraph_format.space_after = Pt(3)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margin(cell, top=90, start=105, bottom=90, end=105):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for tag, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn('w:'+tag))
        if node is None: node = OxmlElement('w:'+tag); mar.append(node)
        node.set(qn('w:w'), str(val)); node.set(qn('w:type'), 'dxa')

def borders(table, color='D9D9D9', size='6'):
    tblPr = table._tbl.tblPr; el = tblPr.first_child_found_in('w:tblBorders')
    if el is None: el = OxmlElement('w:tblBorders'); tblPr.append(el)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement('w:'+edge); tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),size); tag.set(qn('w:color'),color); el.append(tag)

def table(headers, rows, widths=None, font=8.6):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False; borders(t)
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,'294D73'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8.5)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],'F3F7FB')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
                for r in p.runs: r.font.size=Pt(font)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    for row in t.rows:
        for c in row.cells: set_cell_margin(c)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def bullet(text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); p.paragraph_format.space_after=Pt(3); return p

def path(text): doc.add_paragraph(text, style='Code Path')

# Cover
p=doc.add_paragraph(); p.style='Title'; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Patient Record Consultation Record and Doctor Consultation Revision')
p=doc.add_paragraph('End to end technical file inventory and workflow reference')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].font.size=Pt(13); p.runs[0].font.color.rgb=RGBColor(69,88,108)
doc.add_paragraph()
table(['Document item','Value'],[
    ('System','SLSU Clinic Management Information System'),
    ('Repository basis','Current workspace implementation inspected on 15 September 2026'),
    ('Coverage','Blade views, routes, controllers, models, middleware, menu configuration, styles, scripts, database tables, PDF output, and inventory side effects'),
    ('Audience','Developers, reviewers, testers, and deployment personnel'),
], [1.55,5.35], 9.2)
doc.add_paragraph('Purpose', style='Heading 1')
doc.add_paragraph('This document identifies every directly involved component for the revised Patient Record, Consultation Record, and Doctor Consultation features. It also explains how a medical record moves between the nurse, doctor, and nurse completion stages. The current implementation uses the medicalrecord row as the workflow anchor and stores prescription lines in doctor_consultation.')
doc.add_paragraph('Main conclusion', style='Heading 2')
doc.add_paragraph('The three features are not independent pages. They form one campus restricted workflow controlled by record status. A record with status For Doctor appears in the doctor queue. Saving the doctor consultation stores prescription lines, deducts stock on the first doctor pass, and changes the record to For Nurse. Nurse completion changes it to Active. Patient Record exposes the same actions according to the current status.')

doc.add_page_break()
doc.add_paragraph('Workflow Overview', style='Heading 1')
table(['Stage','User action','Route and handler','State or output'],[
 ('Patient search','Search Student Employee or Dependent','POST /patient-view-record → PatientViewMedicalRecordController searchView','Returns encrypted patient identifiers as JSON'),
 ('Patient record','Open selected history','GET /patient-record → records','Displays medicalrecord history and available actions'),
 ('Doctor queue','Open pending intake','GET /doctor-consultations → doctorConsultations','Lists records where status is For Doctor'),
 ('Doctor consultation','Record recommendation and prescription','POST /doctor-consultation/{record} → saveDoctorConsultation','Writes doctor_consultation rows deducts inventory and sets For Nurse'),
 ('Prescription PDF','Preview current form or print saved prescription','POST prescription-pdf or GET saved-prescription-pdf','Streams A5 PDF from doctor-prescription-pdf Blade'),
 ('Consultation record','Review doctor completed cases','GET /consultation-records → consultationRecords','Lists records where status is For Nurse'),
 ('Nurse completion','Complete treatment','POST /nurse-treatment/{record} → saveNurseTreatment','Updates purpose and recommendation then sets Active'),
], [1.15,1.65,2.55,1.55], 8.2)
doc.add_paragraph('Status Transition Rules', style='Heading 2')
table(['Current status','Visible location','Allowed next action','Next status'],[
 ('For Doctor','Doctor Consultations and Patient Record','Save doctor consultation','For Nurse'),
 ('For Nurse','Consultation Record and Patient Record','Complete nurse treatment','Active'),
 ('Active','Patient Record and patient status functions','Normal record maintenance or logout','InActive when logged out'),
], [1.2,2.0,2.2,1.5], 8.7)
doc.add_paragraph('Access Controls', style='Heading 2')
bullet('All routes in scope are inside the authuser middleware group. The middleware requires a session role and redirects unauthenticated sessions to login.')
bullet('Doctor consultation, consultation record, prescription, and nurse treatment routes are additionally inside campus.one. CampusOneOnly returns HTTP 403 unless session campus equals 1.')
bullet('MenuServiceProvider removes doctor-consultations and consultation-records menu items outside Campus 1. This hides the links while middleware still enforces access at the route level.')

doc.add_page_break()
doc.add_paragraph('Patient Record Components', style='Heading 1')
doc.add_paragraph('Patient Record provides patient search, medical history display, edit and soft delete actions, printable history, and status based links into doctor consultation or nurse completion.')
table(['Layer','File','Responsibility'],[
 ('Blade','resources/views/pages/patient-view-record.blade.php','Search page for Student Employee and Dependent records; posts search data and redirects with encrypted ID and role.'),
 ('Blade','resources/views/pages/patient-record.blade.php','History table; print edit delete doctor consultation nurse completion and saved prescription actions.'),
 ('Blade partial','resources/views/modal/editMedical-Record.blade.php','Edit form posted to /update-record including purpose findings vitals recommendation and OTC medicine rows.'),
 ('Blade partial','resources/views/modal/addMedical-Record.blade.php','Included by the record page as the add record dialog.'),
 ('Controller','app/Http/Controllers/PatientViewMedicalRecordController.php','Searches patients loads history edits records synchronizes inventory soft deletes records and generates PDF.'),
 ('Model','app/Medical.php','Eloquent model for medicalrecord with SoftDeletes and the main record fields.'),
 ('Models','app/Student.php and app/Employee.php','Resolve patient identity for Student and Employee roles.'),
 ('Database query','dependent_info table','Resolves Dependent patient identity without a dedicated model in this flow.'),
 ('Inventory','app/Inventory.php and app/Stocks.php','Used by record editing to reconcile OTC quantities and recalculate running stock.'),
], [1.0,2.45,3.45], 8.25)
doc.add_paragraph('Routes and Controller Methods', style='Heading 2')
table(['Method and URI','Controller method','Purpose'],[
 ('GET /patient-view-record','index','Render the search screen.'),
 ('POST /patient-view-record','searchView','Search by role and campus; return JSON search results.'),
 ('GET /patient-record','records','Decrypt patient ID, load history, detect saved prescriptions, and render record page.'),
 ('POST /viewModal','viewModal','Return one medical record with patient information for the edit modal.'),
 ('POST /update-record','updateViewModal','Update medicalrecord and reconcile related OTC inventory inside a transaction.'),
 ('POST /delete-record','delete','Soft delete by setting deleted_at.'),
 ('GET /generate-pdf','generatePDF','Generate the printable patient record output.'),
 ('POST /searchItems','PatientMedicalRecordController searchItems','Inventory autocomplete used by medicine fields.'),
], [2.05,1.9,2.95], 8.35)
doc.add_paragraph('Important Runtime Behavior', style='Heading 2')
bullet('records decrypts the incoming id with AESCipher and restricts the patient and record queries by session campus. Legacy rows with blank or null campus remain visible in selected queries.')
bullet('The record page queries doctor_consultation to decide whether the Print Prescription action should appear.')
bullet('Editing OTC medicine usage calls syncMedicalInventory, findMedicalInventoryRowByMedicine, and recalculateInventoryFrom. These methods update the inventory ledger and the current stock total.')
bullet('The page uses jQuery AJAX, DataTables, SweetAlert2, JSZip, and pdfmake assets. CSRF is taken from the page meta tag.')

doc.add_paragraph('Doctor Consultations Components', style='Heading 1')
doc.add_paragraph('Doctor Consultations consists of the doctor queue, the three step consultation form, prescription generation, stock deduction, and persistence of prescription lines.')
table(['Layer','File','Responsibility'],[
 ('Blade queue','resources/views/pages/doctor-consultations.blade.php','Searchable 12 item paginated queue of records with status For Doctor.'),
 ('Blade form','resources/views/pages/doctor-consultation.blade.php','Three step recommendation prescription and preview interface; medicine stock search and AJAX save.'),
 ('Blade PDF','resources/views/pages/doctor-prescription-pdf.blade.php','A5 prescription layout used for preview and saved prescription printing.'),
 ('CSS','public/css/pages/doctor-consultation.css','Dedicated styling for the consultation workflow and prescription preview.'),
 ('Controller','app/Http/Controllers/PatientMedicalRecordController.php','Loads queue and form validates input persists consultation generates PDFs handles inventory and moves workflow status.'),
 ('Database','doctor_consultation.sql','Table definition and deployment dump for prescription line records.'),
 ('Inventory','app/Stocks.php and app/Inventory.php','Lock stock rows verify quantities deduct stock and append inventory ledger entries.'),
 ('Doctor data','app/Doctor.php and doctors table','Find the active physician for the current campus for the form and PDF.'),
], [1.0,2.55,3.35], 8.2)
doc.add_paragraph('Routes and Controller Methods', style='Heading 2')
table(['Method and URI','Named route','Controller method'],[
 ('GET /doctor-consultations','medical.doctor-consultations','doctorConsultations'),
 ('GET /doctor-consultation/{record}','medical.doctor-consultation','doctorConsultation'),
 ('POST /doctor-consultation/{record}','medical.doctor-consultation.save','saveDoctorConsultation'),
 ('POST /doctor-consultation/{record}/prescription-pdf','medical.doctor-prescription.pdf','prescriptionPdf'),
 ('GET /doctor-consultation/{record}/saved-prescription-pdf','medical.doctor-prescription.saved-pdf','savedPrescriptionPdf'),
 ('POST /searchItems','none','searchItems'),
], [2.85,2.3,1.75], 8.2)
doc.add_paragraph('Save Transaction', style='Heading 2')
table(['Step','Implementation'],[
 ('1 Validate','Recommendation is required. Medicine quantity must be at least 1. Dose unit is limited to the configured list.'),
 ('2 Resolve record','Load Medical by route record ID and current campus or fail with 404.'),
 ('3 Build rows','Create one doctor_consultation row per nonblank medicine and prepare matching inventory deductions.'),
 ('4 Lock stock','On the first For Doctor save lock Stock and latest Inventory rows and reject missing or insufficient stock.'),
 ('5 Deduct inventory','Update stock item_quantity and insert an inventory ledger row with patient stock lot and campus data.'),
 ('6 Replace prescription','Soft delete existing active doctor_consultation rows for the medical record and insert the new rows.'),
 ('7 Move workflow','Append OTC Medicine to purpose when needed and update medicalrecord recommendation and status to For Nurse.'),
], [.55,6.35], 8.6)
doc.add_paragraph('Client Side Dependencies', style='Heading 2')
bullet('The form searches /searchItems as the user types, records stock ID and lot number, shows low stock and expiration warnings, and limits the entered quantity to available stock.')
bullet('The Save button submits FormData with Accept application/json and X Requested With XMLHttpRequest. Validation or stock errors are shown through SweetAlert2.')
bullet('The Print Prescription button uses its own formaction and posts the current unsaved form to the PDF endpoint. Saved prescriptions use the separate GET endpoint from Patient Record.')

doc.add_paragraph('Consultation Record Components', style='Heading 1')
doc.add_paragraph('Consultation Record is the nurse facing list of doctor completed consultations. It combines each medical record with patient identity and a grouped medicine summary, then opens Nurse Treatment for completion.')
table(['Layer','File','Responsibility'],[
 ('Blade','resources/views/pages/consultation-records.blade.php','Displays patient date findings recommendation prescribed medicines and For Nurse status; links to Complete Treatment.'),
 ('Blade','resources/views/pages/nurse-treatment.blade.php','Loads the selected record and saved doctor_consultation rows and submits nurse completion.'),
 ('Controller','app/Http/Controllers/PatientMedicalRecordController.php','consultationRecords nurseTreatment and saveNurseTreatment implement the list and completion flow.'),
 ('Data','medicalrecord and doctor_consultation','medicalrecord is the workflow anchor; doctor_consultation is grouped into a medicine summary by medical record ID.'),
 ('Navigation','resources/data/menus/nurse-menu.json and nurse-attendant-menu.json','Expose Consultation Record for nursing roles.'),
], [1.0,2.7,3.2], 8.3)
doc.add_paragraph('Routes and Controller Methods', style='Heading 2')
table(['Method and URI','Named route','Behavior'],[
 ('GET /consultation-records','medical.consultation-records','Search and paginate status For Nurse records.'),
 ('GET /nurse-treatment/{record}','medical.nurse-treatment','Load Medical patient and active doctor_consultation rows.'),
 ('POST /nurse-treatment/{record}','medical.nurse-treatment.save','Validate purpose and recommendation then set status Active.'),
], [2.45,2.25,2.2], 8.5)
doc.add_paragraph('Query Details', style='Heading 2')
bullet('consultationRecords groups medicine_name and quantity with GROUP_CONCAT from active doctor_consultation rows.')
bullet('Patient names are resolved with role conditional joins to student_info, employee_info, and dependent_info, then combined with COALESCE.')
bullet('Only records with status For Nurse and no medicalrecord deleted_at value are returned. Results are ordered by date and time and paginated 12 per page.')
bullet('Search matches patient ID plus first or last name across all three patient types.')

doc.add_paragraph('Shared Supporting Components', style='Heading 1')
table(['Component','Path','Why it is involved'],[
 ('Routes','routes/web.php','Declares every active endpoint and middleware grouping.'),
 ('Route middleware aliases','app/Http/Kernel.php','Maps authuser and campus.one aliases.'),
 ('Authentication middleware','app/Http/Middleware/authUser.php','Requires a session role before allowing the scoped pages.'),
 ('Campus restriction','app/Http/Middleware/CampusOneOnly.php','Restricts the consultation workflow to Campus 1.'),
 ('Menu selection','app/Providers/MenuServiceProvider.php','Loads role specific menus and removes campus restricted entries elsewhere.'),
 ('Layout','resources/views/layouts/contentLayoutMaster.blade.php','Base layout extended by all three principal screens.'),
 ('Menu data','resources/data/menus/doctor-menu.json','Links Doctor users to Consultations and Medical Record.'),
 ('Menu data','resources/data/menus/nurse-menu.json','Links Nurse users to Patient Record and Consultation Record.'),
 ('Menu data','resources/data/menus/nurse-attendant-menu.json','Provides the corresponding nurse attendant links.'),
 ('Assets','public/vendors and public/js/scripts/datatables/datatable.js','DataTables and common client side table behavior used by Patient Record and Doctor Consultation.'),
 ('PDF package','barryvdh/laravel-dompdf via composer.json','Renders prescription and patient record PDF views.'),
], [1.35,2.75,2.8], 8.0)
doc.add_paragraph('Database Objects and Key Relationships', style='Heading 2')
table(['Object','Key fields used','Relationship'],[
 ('medicalrecord','id patientId role campus status purpose findings recommendation vitals deleted_at','One row is the workflow anchor. id is referenced by doctor_consultation.patientId.'),
 ('doctor_consultation','patientId quantity medicine_name dose route frequency duration instruction deleted_at','Multiple prescription lines belong to one medicalrecord row. The field name patientId actually stores medicalrecord.id.'),
 ('student_info','StudentNo name campus','Joined when medicalrecord.role is Student.'),
 ('employee_info','id name campus','Joined when role is Employee.'),
 ('dependent_info','id name campus services','Joined when role is Dependent.'),
 ('doctors','campus specialization deleted_at','Selects the first active physician for the current campus.'),
 ('stock','id item_name item_quantity lotno campus deleted_at','Authoritative current stock record used for prescription deductions.'),
 ('inventory','stockId patientId lotno item_stock stock_less remaining_stock campus date','Append only style stock movement ledger and record edit reconciliation source.'),
], [1.25,2.75,2.9], 8.0)

doc.add_paragraph('Revision and Deployment Checklist', style='Heading 1')
doc.add_paragraph('Files that must move together', style='Heading 2')
for item in [
 'routes/web.php',
 'app/Http/Controllers/PatientMedicalRecordController.php',
 'app/Http/Controllers/PatientViewMedicalRecordController.php',
 'app/Http/Middleware/CampusOneOnly.php and app/Http/Kernel.php',
 'app/Providers/MenuServiceProvider.php',
 'resources/views/pages/patient-view-record.blade.php and patient-record.blade.php',
 'resources/views/modal/editMedical-Record.blade.php and addMedical-Record.blade.php',
 'resources/views/pages/doctor-consultations.blade.php and doctor-consultation.blade.php',
 'resources/views/pages/consultation-records.blade.php and nurse-treatment.blade.php',
 'resources/views/pages/doctor-prescription-pdf.blade.php',
 'public/css/pages/doctor-consultation.css',
 'resources/data/menus/doctor-menu.json nurse-menu.json nurse-attendant-menu.json and vertical-menu.json',
 'doctor_consultation.sql or an equivalent production migration',
]: bullet(item)
doc.add_paragraph('Verification Checks', style='Heading 2')
table(['Check','Expected result'],[
 ('Authentication','Anonymous users are redirected to login.'),
 ('Campus restriction','A non Campus 1 request to consultation routes returns 403 and menu entries are absent.'),
 ('Doctor queue','Only active nondeleted For Doctor records appear; name and ID search works for all roles.'),
 ('Doctor save','Recommendation validation works; each selected medicine is persisted; stock and inventory remain consistent.'),
 ('Repeat save','Existing prescription rows are soft deleted and replaced; stock is not deducted again after status becomes For Nurse.'),
 ('Prescription output','Current form PDF and saved prescription PDF render correct patient doctor medicine and instruction data.'),
 ('Consultation list','For Nurse records display grouped medicine names and quantities and open Nurse Treatment.'),
 ('Nurse completion','Purpose and recommendation save and status becomes Active.'),
 ('Patient record','Edit delete print consultation completion and saved prescription actions appear under the correct conditions.'),
 ('Cross campus data','Every record stock doctor and patient lookup is checked for intended campus behavior including legacy null campus rows.'),
], [1.55,5.35], 8.35)
doc.add_paragraph('Implementation Notes Requiring Attention', style='Heading 2')
bullet('doctor_consultation.sql is a database dump with DROP TABLE and sample data, not a Laravel migration. Production deployment should use a reviewed migration or controlled import procedure.')
bullet('doctor_consultation.patientId is named like a patient identifier but stores medicalrecord.id. This should be documented consistently or renamed through a migration to reduce join errors.')
bullet('The SQL schema limits instruction and several prescription fields to varchar 45. The controller accepts unrestricted strings, so long input may be truncated or rejected depending on SQL mode.')
bullet('The employee joins use employee_info.id in the revised consultation flow. Other older patient record code sometimes references AgencyNumber. Confirm existing medicalrecord.patientId values follow the revised convention.')
bullet('Legacy copies of web.php and controllers exist under public data or other nonstandard locations. The active Laravel files are routes/web.php and app/Http/Controllers. Do not deploy the copies as routing or controller sources.')

# Footer and page numbers
for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('SLSU Clinic Management Information System  |  Revision Reference  |  ')
    run.font.size=Pt(8); run.font.color.rgb=RGBColor(92,101,112)
    fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
    instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text='PAGE'
    fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)

# Keep headings with following content and repeat table headers.
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next = True
for t in doc.tables:
    trPr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)

doc.core_properties.title = 'Patient Record Consultation Record and Doctor Consultation Revision'
doc.core_properties.subject = 'End to end technical file inventory and workflow reference'
doc.core_properties.author = 'SLSU Clinic Management Information System Project'
doc.save(OUT)
print(OUT)
